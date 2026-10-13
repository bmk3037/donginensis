"""가상 예시 프로젝트(projects/예시)를 만든다 — 도구 시험·사용법 확인용. 모든 값은 가상이다.

  python3 example/make_sample.py && python3 grant.py fill 예시 && python3 grant.py build 예시
"""
import shutil
import sys
import warnings
from pathlib import Path

warnings.filterwarnings('ignore')
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import docx  # noqa: E402
from docx.shared import Pt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from grantkit.project import FONTS, create  # noqa: E402

P = create('예시')


def plan_docx():
    d = docx.Document()
    d.add_heading('2026년 예시 성장사다리 지원사업 사업계획서', 0)
    d.add_paragraph('※ 도구 시험용 가상 문서입니다.')
    t = d.add_table(rows=6, cols=2)
    t.style = 'Table Grid'
    for i, (k, v) in enumerate([('기업명', '(주)예시기업'), ('대표자', '김예시'), ('사업자등록번호', '000-00-00000'),
                                ('과제명', '스마트 공정 모니터링 시스템 구축'), ('사업기간', '2026. 5. 1. ~ 2026. 11. 30.'),
                                ('사업비', '총 30,000,000원 (정부지원금 21,000,000원, 자부담 현금 9,000,000원)')]):
        t.cell(i, 0).text, t.cell(i, 1).text = k, v
    d.add_heading('1. 세부 추진내용', 1)
    d.add_paragraph('T1. 공정 설비 데이터 수집 장치 설치: 생산라인 3개 라인에 전력·진동 센서와 데이터 수집 게이트웨이를 설치한다.')
    d.add_paragraph('T2. 모니터링 대시보드 구축: 수집 데이터를 실시간으로 보여주는 웹 대시보드와 이상 알림 기능을 개발한다.')
    d.add_heading('2. 성과지표', 1)
    t2 = d.add_table(rows=3, cols=3)
    t2.style = 'Table Grid'
    for i, row in enumerate([('지표', '목표', '증빙'), ('설비 데이터 수집 라인 수', '3개 라인', '설치 사진, 수집 화면'),
                             ('설비 비가동 시간 감소율', '10%', '전후 비교 데이터')]):
        for j, v in enumerate(row):
            t2.cell(i, j).text = v
    d.save(P / 'input/plan/사업계획서.docx')


def forms():
    from hwpx import HwpxDocument
    doc = HwpxDocument.new()
    doc.paragraphs[0].text = '대 금 청 구 서'
    t = doc.add_table(6, 4)
    labels = [('사업명', 0, 0), ('과제명', 1, 0), ('기업명', 2, 0), ('대표자', 2, 2), ('청구금액', 3, 0),
              ('사업기간', 4, 0), ('입금계좌', 5, 0)]
    for lab, r, c in labels:
        t.cell(r, c).text = lab
    doc.add_paragraph('위와 같이 지원사업 대금을 청구합니다.')
    doc.add_paragraph('2026년    월    일')
    doc.add_paragraph('청구인:                (인)')
    doc.save_to_path(str(P / 'input/forms/대금청구서.hwp'))

    d = docx.Document()
    d.add_heading('수행 결과 확인서', 1)
    t = d.add_table(rows=4, cols=2)
    t.style = 'Table Grid'
    for i, k in enumerate(['과제명', '수행기간', '수행 결과 요약', '현장 사진']):
        t.cell(i, 0).text = k
    d.add_paragraph('작성일:   년   월   일')
    d.save(P / 'input/forms/수행결과확인서.docx')

    wb = Workbook()
    ws = wb.active
    ws.title = '집행내역'
    ws['A1'] = '사업비 집행내역서'
    ws['A3'], ws['C3'] = '기업명', '사업명'
    ws.merge_cells('D3:F3')
    ws.append([])
    ws.append(['비목', '계획', '집행', '잔액'])
    for col, w in zip('ABCDEF', (14, 22, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    wb.save(P / 'input/forms/집행내역서.xlsx')


def photos():
    font = ImageFont.truetype(str(FONTS / 'Pretendard-Bold.otf'), 46)
    shots = [('E01', '라인1_센서설치', (40, 90, 140), '2026:07:14 10:22:00'),
             ('E01', '라인2_센서설치', (40, 110, 120), '2026:07:15 14:05:00'),
             ('E01', '라인3_게이트웨이', (60, 80, 150), '2026:07:16 09:40:00'),
             ('E02', '대시보드_화면', (20, 40, 70), '2026:09:02 16:30:00')]
    for eid, name, color, when in shots:
        im = Image.new('RGB', (1600, 1200), color)
        dr = ImageDraw.Draw(im)
        dr.text((80, 520), f'샘플 사진 · {name}\n(도구 시험용, 실제 증빙 아님)', fill='white', font=font)
        exif = Image.Exif()
        exif[0x8769] = {36867: when}
        exif[306] = when
        im.save(P / f'evidence/{eid}_{name}.jpg', exif=exif)
    (P / 'evidence/E03_시험성적서.pdf').write_bytes(b'%PDF-1.4\n% sample\n')


plan_docx()
forms()
photos()
for sub in ('fill', 'docs'):
    for f in (HERE / sub).glob('*.yaml'):
        shutil.copy(f, P / sub / f.name)
shutil.copy(HERE / 'project.yaml', P / 'project.yaml')
print('예시 프로젝트:', P)
