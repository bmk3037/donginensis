"""HWP·HWPX 양식 → DOCX 변환 (문단·표·병합·칸 너비·배경색·글자 크기/굵기·정렬·그림).

기관이 한글 양식만 줄 때 워드로 바꿔 채우기 위한 것이다. 쪽 배치·도형·글상자·각주 같은 세부 서식은
옮기지 않으므로, 변환본은 한 번 열어 보고 필요하면 손본다.
"""
from __future__ import annotations

import io
import tempfile
import warnings
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

FONT = '맑은 고딕'
HWPUNIT_PER_MM = 7200 / 25.4
# 한글의 양쪽·배분 정렬은 워드에서 짧은 칸 글자를 벌려 놓으므로 왼쪽 정렬로 옮긴다
ALIGN = {'LEFT': WD_ALIGN_PARAGRAPH.LEFT, 'RIGHT': WD_ALIGN_PARAGRAPH.RIGHT, 'CENTER': WD_ALIGN_PARAGRAPH.CENTER}


def _local(e):
    return etree.QName(e).localname


def _child(e, name):
    for c in e:
        if _local(c) == name:
            return c
    return None


class _Header:
    """header.xml 의 글자·문단·테두리 모양 표"""

    def __init__(self, xml: bytes | None):
        self.char, self.para, self.fill = {}, {}, {}
        if not xml:
            return
        root = etree.fromstring(xml)
        for e in root.iter():
            tag = _local(e)
            if tag == 'charPr':
                h = int(e.get('height', '1000'))
                color = e.get('textColor', '#000000')
                self.char[e.get('id')] = {'size': h / 100, 'bold': _child(e, 'bold') is not None,
                                          'color': None if color.upper() in ('#000000', 'NONE') else color.lstrip('#')}
            elif tag == 'paraPr':
                a = _child(e, 'align')
                self.para[e.get('id')] = a.get('horizontal', 'JUSTIFY') if a is not None else 'JUSTIFY'
            elif tag == 'borderFill':
                for f in e.iter():
                    if _local(f) == 'winBrush':
                        fc = f.get('faceColor', 'none')
                        if fc and fc.lower() != 'none' and fc.upper() != '#FFFFFF':
                            self.fill[e.get('id')] = fc.lstrip('#')


def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color.upper())


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    tcpr.append(sh)


def _borders(table):
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), '000000')
        b.append(e)
    table._tbl.tblPr.append(b)


class Converter:
    def __init__(self, src: Path):
        self.src = src
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            from hwpx import HwpxDocument
            doc = HwpxDocument.open(str(src))
            buf = io.BytesIO()
            with tempfile.NamedTemporaryFile(suffix='.hwpx') as tmp:   # .hwp → .hwpx package
                doc.save_to_path(tmp.name)
                buf.write(Path(tmp.name).read_bytes())
        self.zip = zipfile.ZipFile(buf)
        names = self.zip.namelist()
        self.header = _Header(self.zip.read('Contents/header.xml') if 'Contents/header.xml' in names else None)
        self.sections = sorted(n for n in names if n.startswith('Contents/section') and n.endswith('.xml'))
        self.bin = {Path(n).stem: n for n in names if n.startswith('BinData/')}
        self.out = Document()
        st = self.out.styles['Normal']
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        st.paragraph_format.space_after = Pt(0)
        self.content_mm = 170

    # -- page
    def _page(self, sec_root):
        for e in sec_root.iter():
            if _local(e) == 'pagePr':
                w, h = int(e.get('width', 59528)), int(e.get('height', 84186))
                m = _child(e, 'margin')
                s = self.out.sections[0]
                s.page_width, s.page_height = Mm(w / HWPUNIT_PER_MM), Mm(h / HWPUNIT_PER_MM)
                if m is not None:
                    s.left_margin = Mm(int(m.get('left', 8504)) / HWPUNIT_PER_MM)
                    s.right_margin = Mm(int(m.get('right', 8504)) / HWPUNIT_PER_MM)
                    s.top_margin = Mm((int(m.get('top', 5668)) + int(m.get('header', 4252))) / HWPUNIT_PER_MM)
                    s.bottom_margin = Mm((int(m.get('bottom', 4252)) + int(m.get('footer', 4252))) / HWPUNIT_PER_MM)
                self.content_mm = (s.page_width - s.left_margin - s.right_margin) / 36000
                return

    # -- paragraphs
    def _emit_paragraph(self, container, p_el, first_para=None):
        """hp:p → docx paragraph(s); tables/pictures inside runs are emitted in order."""
        para = first_para if first_para is not None else container.add_paragraph()
        align = self.header.para.get(p_el.get('paraPrIDRef'))
        if align in ALIGN:
            para.alignment = ALIGN[align]
        used = False
        for run in p_el:
            if _local(run) != 'run':
                continue
            cp = self.header.char.get(run.get('charPrIDRef'), {})
            for item in run:
                tag = _local(item)
                if tag == 't':
                    text = ''.join(item.itertext())
                    if text:
                        _font(para.add_run(text), cp.get('size'), cp.get('bold'), cp.get('color'))
                        used = True
                elif tag in ('lineBreak',):
                    para.add_run().add_break()
                elif tag == 'tab':
                    para.add_run('\t')
                elif tag == 'tbl':
                    if hasattr(container, 'add_table') and not hasattr(container, '_tc'):
                        self._emit_table(item)
                        para = container.add_paragraph()
                    else:                       # nested table inside a cell → flatten as text lines
                        for sub in item.iter():
                            if _local(sub) == 'p':
                                t = ''.join(x.text or '' for x in sub.iter() if _local(x) == 't')
                                if t.strip():
                                    _font(container.add_paragraph().add_run(t), 9)
                elif tag == 'pic':
                    self._emit_picture(para, item)
                    used = True
                elif tag in ('ctrl', 'secPr', 'colPr'):
                    continue
        return para, used

    def _emit_picture(self, para, pic):
        ref, w = None, None
        for e in pic.iter():
            n = _local(e)
            if n == 'img':
                ref = e.get('binaryItemIDRef')
            elif n == 'curSz' and w is None:
                w = int(e.get('width', 0))
        if not ref or ref not in self.bin:
            return
        data = self.zip.read(self.bin[ref])
        width = Mm(min(w / HWPUNIT_PER_MM, self.content_mm) if w else 60)
        try:
            para.add_run().add_picture(io.BytesIO(data), width=width)
        except Exception:
            pass                                    # unsupported image format (e.g. WMF)

    # -- tables
    def _emit_table(self, tbl):
        cells = []
        for tr in tbl:
            if _local(tr) != 'tr':
                continue
            for tc in tr:
                if _local(tc) != 'tc':
                    continue
                addr, span, sz = _child(tc, 'cellAddr'), _child(tc, 'cellSpan'), _child(tc, 'cellSz')
                cells.append({'r': int(addr.get('rowAddr')), 'c': int(addr.get('colAddr')),
                              'rs': int(span.get('rowSpan', 1)) if span is not None else 1,
                              'cs': int(span.get('colSpan', 1)) if span is not None else 1,
                              'w': int(sz.get('width', 0)) if sz is not None else 0,
                              'fill': self.header.fill.get(tc.get('borderFillIDRef')),
                              'el': tc})
        if not cells:
            return
        nrows = max(c['r'] + c['rs'] for c in cells)
        ncols = max(c['c'] + c['cs'] for c in cells)
        widths = [0] * ncols
        for c in cells:                              # widths from single-column cells first
            if c['cs'] == 1:
                widths[c['c']] = max(widths[c['c']], c['w'])
        for c in cells:
            if c['cs'] > 1:
                known = sum(widths[c['c']:c['c'] + c['cs']])
                missing = [i for i in range(c['c'], c['c'] + c['cs']) if not widths[i]]
                for i in missing:
                    widths[i] = max(0, (c['w'] - known)) // len(missing) or 1
        total = sum(widths) or 1
        scale = self.content_mm / (total / HWPUNIT_PER_MM) if total / HWPUNIT_PER_MM > self.content_mm else 1

        t = self.out.add_table(rows=nrows, cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _borders(t)
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Mm(w / HWPUNIT_PER_MM * scale)
        for c in cells:
            cell = t.cell(c['r'], c['c'])
            if c['rs'] > 1 or c['cs'] > 1:
                cell = cell.merge(t.cell(c['r'] + c['rs'] - 1, c['c'] + c['cs'] - 1))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if c['fill']:
                _shade(cell, c['fill'])
            sub = _child(c['el'], 'subList')
            first = cell.paragraphs[0]
            if sub is not None:
                started = False
                for p_el in sub:
                    if _local(p_el) != 'p':
                        continue
                    self._emit_paragraph(cell, p_el, first_para=None if started else first)
                    started = True

    def convert(self, out: Path) -> Path:
        first = True
        for name in self.sections:
            root = etree.fromstring(self.zip.read(name))
            if first:
                self._page(root)
                first = False
            for p_el in root:
                if _local(p_el) == 'p':
                    self._emit_paragraph(self.out, p_el)
        self.out.save(str(out))
        return out


def convert(src: Path, out: Path) -> Path:
    return Converter(src).convert(out)
