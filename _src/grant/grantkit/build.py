"""자체 작성 문서(결과보고서·컨설팅보고서 등)와 사진대지를 동인엔시스 문서 서식 DOCX(+PDF)로 만든다.

서식 틀: assets/dongin_template.docx — 통합경영 매뉴얼(files/iso/DI-IMS-M-01_Rev00.docx)에서 본문만 비운 것.
스타일(제목 1 빨간 밑줄·제목 2·글머리표), 머리말·꼬리말, 용지·여백을 그대로 쓴다.

문서 명세(docs/*.yaml):
  출력: 04_결과보고서
  문서명: 결과보고서                  # 머리말·꼬리말에 들어감
  문서번호: ""                        # 있으면 머리말·꼬리말에 함께 표기 (지어내지 않음)
  표지:
    제목: "{{사업.과제명}}"
    부제: [결과보고서, "{{사업.사업명}}"]
    항목: {사업명: "...", 수행기업: "...", 사업기간: "..."}
    결재: {작성: 홍길동 과장, 검토: ..., 승인: ... 대표이사}   # 없으면 결재란 생략
    날짜: "{{청구.청구일|날짜:%Y년 %m월 %d일}}"
  본문:
    - 제목: 1. 사업 개요              # 제목 1 (파란 굵은 글씨 + 빨간 밑줄)
    - 소제목: 가. 추진 목적           # 제목 2
    - 문단: ...
    - 목록: [..., ...]               # 매뉴얼과 같은 글머리표
    - 표: {머리: [구분, 계획, 실적], 행: [[...]], 너비: [30, 60, 60], 합계행: false, 주석: ...}
    - 사진: [{파일: "evidence/E01_*", 설명: 설치 완료 정면}]   # 한 줄에 2장 (열: 3 으로 변경 가능)
    - 강조: 미완료 항목은 12월 중 완료 예정   # 매뉴얼의 노란 강조 문장
    - 쪽나눔: true
  사진대지: true                      # 끝에 증빙 사진 전체를 사진대지로 붙임
"""
from __future__ import annotations

import io
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips

from . import media
from .project import Project, fmt_date, render

ASSETS = Path(__file__).resolve().parent / 'assets'
TEMPLATE = ASSETS / 'dongin_template.docx'
LOGO = ASSETS / 'logo.png'
MANUAL = Path(__file__).resolve().parents[3] / 'files' / 'iso' / 'DI-IMS-M-01_Rev00.docx'
COMPANY = '동인엔시스'

FONT = 'Malgun Gothic'
INK = RGBColor(0x26, 0x34, 0x45)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x65, 0x77)
HEAD_FILL, ZEBRA_FILL, LINE = '1857A5', 'FFFFFF', 'D3DDE8'   # 동인엔시스 서식: 연한 파란 칸 없음(머리행만 파랑, 대표 지시 2026.10.23)
BULLET_NUM_ID = '2'          # 매뉴얼의 글머리표 목록 번호


def ensure_template():
    """서식 틀이 없으면 매뉴얼에서 다시 만든다."""
    if TEMPLATE.exists():
        return
    d = Document(str(MANUAL))
    body = d.element.body
    for el in list(body):
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    ASSETS.mkdir(exist_ok=True)
    d.save(str(TEMPLATE))


# ---- low-level helpers -----------------------------------------------------------------------

def _font(run, size=10.5, bold=None, color=INK):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _spacing(p, after=100, line=276, before=None):
    pf = p.paragraph_format
    pf.space_after = Twips(after)
    if before is not None:
        pf.space_before = Twips(before)
    pf.line_spacing = line / 240


def _cell_style(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '6')
        e.set(qn('w:color'), LINE)
        borders.append(e)
    tcpr.append(borders)
    mar = OxmlElement('w:tcMar')
    for side, v in (('top', 60), ('left', 90), ('bottom', 60), ('right', 90)):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcpr.append(mar)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


NUM_HEADS = {'번호', '순번', '연번', 'No', 'No.', 'NO', '구분', '순서'}


def _num_col(head, j):
    """머리글이 번호·순번·연번·No·구분인 칸"""
    return bool(head) and j < len(head) and str(head[j]).strip() in NUM_HEADS


def _is_marker(v):
    """①② 같은 원 숫자, (1)·1) 같은 짧은 번호 표시"""
    import re
    return bool(re.fullmatch(r'[①-⑳]|\(?\d{1,2}\)', str('' if v is None else v).strip()))


def _cell_text(cell, text, head=False, align=None, bold=None):
    cell.text = ''
    for i, line in enumerate(str('' if text is None else text).split('\n')):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if head else align
        _font(p.add_run(line), 9.5, True if head else bold, WHITE if head else INK)


def _row_flags(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    if header:
        trpr.append(OxmlElement('w:tblHeader'))


def _is_number(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    s = str(v).replace(',', '').replace('%', '').replace('원', '').strip()
    return bool(s) and s.lstrip('-').replace('.', '', 1).isdigit()


class Writer:
    def __init__(self, title: str, doc_no: str = ''):
        ensure_template()
        self.doc = Document(str(TEMPLATE))
        self.photo_no = 0
        s = self.doc.sections[0]
        self.content_mm = (s.page_width - s.left_margin - s.right_margin) / 36000
        self._header_footer(title, doc_no)

    def _header_footer(self, title, doc_no):
        s = self.doc.sections[0]
        label = f'{title} | {doc_no}' if doc_no else title
        hp = s.header.paragraphs[0]
        if hp.runs:
            hp.runs[0].text = f' {label}'
            for r in hp.runs[1:]:
                r.text = ''
        else:
            _font(hp.add_run(label), 9, color=None)
        first = s.footer.paragraphs[0]._p.find(qn('w:r'))     # "동인엔시스 | … | 페이지 " + PAGE / 전체 NUMPAGES
        if first is not None and first.find(qn('w:t')) is not None:
            first.find(qn('w:t')).text = f'{COMPANY} | {label} | 페이지 '

    # -- cover (매뉴얼 표지와 같은 배치)
    def cover(self, c: dict):
        p = self.doc.add_paragraph()
        _spacing(p, before=2200)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if LOGO.exists():
            p.add_run().add_picture(str(LOGO), width=Mm(95.25))
        p = self.doc.add_paragraph(style='Title')
        _spacing(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(str(c.get('제목', ''))), 20, True)
        rule = self.doc.add_paragraph()
        _spacing(rule)
        ppr = rule._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '24'), ('w:space', '1'), ('w:color', 'E83E30')):
            b.set(qn(k), v)
        bdr.append(b)
        ppr.insert(0, bdr)
        rule.paragraph_format.left_indent = Twips(2800)
        rule.paragraph_format.right_indent = Twips(2800)
        subs = c.get('부제') or []
        for line in ([subs] if isinstance(subs, str) else subs):
            p = self.doc.add_paragraph()
            _spacing(p)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(str(line)), 10.5)
        _spacing(self.doc.add_paragraph())
        items = c.get('항목') or {}
        if items:
            self.table({'머리': ['항목', '내용'], '행': [[k, v] for k, v in items.items()], '너비': [40, 130]})
        appr = c.get('결재')
        if appr:
            roles = list(appr.keys())
            self.table({'머리': ['구분', *roles],
                        '행': [['담당', *appr.values()], ['서명', *[''] * len(roles)], ['일자', *[''] * len(roles)]],
                        '너비': [25] + [145 / len(roles)] * len(roles), '가운데': True, '줄무늬': False})
        if c.get('날짜') or c.get('기관'):
            for _ in range(3):
                self.doc.add_paragraph()
            for key, size, bold in (('날짜', 11, False), ('기관', 14, True)):
                if c.get(key):
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _font(p.add_run(str(c[key])), size, bold)
        self.page_break()

    # -- blocks
    def heading(self, text, level=1):
        p = self.doc.add_paragraph(style='Heading 1' if level == 1 else 'Heading 2')
        p.add_run(str(text))
        p.paragraph_format.keep_with_next = True

    def para(self, text):
        for chunk in str(text).split('\n\n'):
            p = self.doc.add_paragraph()
            _spacing(p)
            _font(p.add_run(chunk.strip()))

    def note(self, text):
        p = self.doc.add_paragraph()
        _spacing(p)
        _font(p.add_run(str(text))).font.highlight_color = WD_COLOR_INDEX.YELLOW

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style='List Paragraph')
            _spacing(p)
            ppr = p._p.get_or_add_pPr()
            num = OxmlElement('w:numPr')
            ilvl, nid = OxmlElement('w:ilvl'), OxmlElement('w:numId')
            ilvl.set(qn('w:val'), '0')
            nid.set(qn('w:val'), BULLET_NUM_ID)
            num.append(ilvl)
            num.append(nid)
            ppr.insert(1 if ppr.find(qn('w:pStyle')) is not None else 0, num)
            _font(p.add_run(str(it)), color=None)

    def table(self, spec: dict):
        head = spec.get('머리') or []
        rows = spec.get('행') or []
        ncol = len(head) or (len(rows[0]) if rows else 0)
        if not ncol:
            return
        t = self.doc.add_table(rows=(1 if head else 0) + len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = (spec.get('너비') or [self.content_mm / ncol] * ncol)[:ncol]
        scale = self.content_mm / sum(widths)
        r0 = 0
        if head:
            for j, h in enumerate(head):
                c = t.cell(0, j)
                _cell_style(c, HEAD_FILL)
                _cell_text(c, h, head=True)
            _row_flags(t.rows[0], header=True)
            r0 = 1
        last = len(rows) - 1
        for i, row in enumerate(rows):
            shaded = (spec.get('줄무늬', True) and i % 2 == 1) or (spec.get('합계행') and i == last)
            total = spec.get('합계행') and i == last
            for j in range(ncol):
                v = row[j] if j < len(row) else ''
                num = _is_number(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    v = f'{v:,}'
                c = t.cell(r0 + i, j)
                _cell_style(c, ZEBRA_FILL if shaded else 'FFFFFF')
                mid = spec.get('가운데') or _num_col(head, j) or _is_marker(v)   # 번호 칸·①② 표시는 가운데(대표 지시 2026.10.23)
                align = WD_ALIGN_PARAGRAPH.CENTER if mid else (WD_ALIGN_PARAGRAPH.RIGHT if num else None)
                _cell_text(c, v, align=align, bold=True if total else None)
            _row_flags(t.rows[r0 + i])
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Mm(w * scale)
        p = self.doc.add_paragraph()
        _spacing(p, before=60 if spec.get('주석') else None, after=60)
        if spec.get('주석'):
            _font(p.add_run(str(spec['주석'])), 9, color=MUTED)

    def photos(self, items: list[dict], project: Project, per_row: int = 2, show_date=False):
        entries = []
        for it in items:
            src = it.get('파일목록') or project.resolve_files(it['파일'])
            files = [Path(f) for f in src if media.is_image(Path(f))]
            for k, f in enumerate(files):
                desc = it.get('설명') or f.stem.split('_', 1)[-1]
                if len(files) > 1 and it.get('설명'):
                    desc = f'{desc} ({k + 1})'
                entries.append((f, desc))
        if not entries:
            self.note('【사진 없음: ' + ', '.join(str(i.get('파일', i.get('설명', ''))) for i in items) + '】')
            return
        cell_w = self.content_mm / per_row
        for start in range(0, len(entries), per_row):
            chunk = entries[start:start + per_row]
            t = self.doc.add_table(rows=2, cols=per_row)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for j in range(per_row):
                c_img, c_cap = t.cell(0, j), t.cell(1, j)
                c_img.width = c_cap.width = Mm(cell_w)
                _cell_style(c_img, 'FFFFFF')
                _cell_style(c_cap, ZEBRA_FILL)
                if j >= len(chunk):
                    continue
                f, desc = chunk[j]
                data, (w, h) = media.prepare(f)
                inner_w = cell_w - 5
                box_h = inner_w * 0.75
                pw = inner_w if w / h >= inner_w / box_h else box_h * w / h
                p = c_img.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(io.BytesIO(data), width=Mm(pw))
                self.photo_no += 1
                cap = f'[사진 {self.photo_no}] {desc}'
                if show_date:
                    d, why = media.taken_date(f)
                    cap += f'\n촬영일: {fmt_date(d) if d else why}'
                _cell_text(c_cap, cap, align=WD_ALIGN_PARAGRAPH.CENTER)
            for row in t.rows:
                _row_flags(row)
            _spacing(self.doc.add_paragraph(), after=60)

    def page_break(self):
        p = self.doc.add_paragraph()
        _spacing(p)
        p.add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: Path):
        self.doc.save(str(path))


def _finish(w: Writer, out: Path, pdf: bool) -> list[Path]:
    out.parent.mkdir(parents=True, exist_ok=True)
    w.save(out)
    outs = [out]
    if pdf:
        p = media.to_pdf(out, out.parent)
        if p:
            outs.append(p)
    return outs


def build_doc(project: Project, spec_path: Path, pdf: bool = True) -> list[Path]:
    spec = render(yaml.safe_load(spec_path.read_text(encoding='utf-8')) or {}, project.data())
    w = Writer(spec.get('문서명') or spec_path.stem, spec.get('문서번호', ''))
    if spec.get('표지'):
        w.cover(spec['표지'])
    for block in spec.get('본문') or []:
        if not isinstance(block, dict):
            w.para(block)
            continue
        if '제목' in block:
            w.heading(block['제목'], 1)
        if '소제목' in block:
            w.heading(block['소제목'], 2)
        if '문단' in block:
            w.para(block['문단'])
        if '목록' in block:
            w.bullets(block['목록'])
        if '표' in block:
            w.table(block['표'])
        if '사진' in block:
            w.photos(block['사진'], project, int(block.get('열', 2)))
        if '강조' in block:
            w.note(block['강조'])
        if block.get('쪽나눔'):
            w.page_break()
    if spec.get('사진대지'):
        w.page_break()
        w.heading('[별첨] 증빙 사진대지', 1)
        w.photos(evidence_photo_items(project), project, 2, show_date=True)
    return _finish(w, project.path('output', (spec.get('출력') or spec_path.stem) + '.docx'), pdf)


def evidence_photo_items(project: Project) -> list[dict]:
    items = []
    for e in project.data().get('증빙목록') or []:
        imgs = [f for f in project.evidence_files(e['id']) if media.is_image(f)]
        if imgs:
            items.append({'파일목록': imgs, '설명': f"{e['id']} {e.get('항목', '')}".strip()})
    return items


def build_photosheet(project: Project, per_row: int = 2, pdf: bool = True) -> list[Path]:
    data = project.data()
    biz = data.get('사업', {})
    period = biz.get('사업기간') or {}
    w = Writer('증빙 사진대지')
    w.heading('증빙 사진대지', 1)
    w.table({'머리': ['사업명', '기업명', '사업기간'],
             '행': [[biz.get('사업명', ''), data.get('기업', {}).get('기업명', ''),
                    f"{fmt_date(period.get('시작', ''))} ~ {fmt_date(period.get('종료', ''))}"]],
             '너비': [70, 50, 50], '가운데': True})
    w.photos(evidence_photo_items(project), project, per_row, show_date=True)
    return _finish(w, project.path('output', '사진대지.docx'), pdf)
