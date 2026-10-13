"""요청자료 체크리스트(XLSX)와 제출 현황."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .project import Project

# 동인엔시스 문서 서식 (통합경영 매뉴얼·기록양식집과 같은 색)
THIN = Side(style='thin', color='D3DDE8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='1857A5')
ZEBRA = PatternFill('solid', fgColor='F0F5FB')
DONE = PatternFill('solid', fgColor='E3F4EA')
TODO = PatternFill('solid', fgColor='FCE6E3')
INK = '263445'
F = '맑은 고딕'


def evidence_status(project: Project) -> list[dict]:
    rows = []
    for e in project.data().get('증빙목록') or []:
        files = project.evidence_files(e['id'])
        rows.append({**e, '파일': [str(f.relative_to(project.path('evidence'))) for f in files],
                     '제출': bool(files)})
    return rows


def _sheet(ws, title, head, widths, rows, status_col=None):
    ws.append([title])
    ws['A1'].font = Font(name=F, size=14, bold=True, color='1857A5')
    ws.append([])
    ws.append(head)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        c = ws.cell(3, j)
        c.font, c.fill, c.border = Font(name=F, bold=True, color='FFFFFF'), HEAD, BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
        i = ws.max_row
        for j in range(1, len(head) + 1):
            c = ws.cell(i, j)
            c.font, c.border = Font(name=F, size=10, color=INK), BOX
            if (i - 3) % 2 == 0:
                c.fill = ZEBRA
            c.alignment = Alignment(vertical='top', wrap_text=True, horizontal='center' if j in (1, status_col) else None)
        if status_col:
            c = ws.cell(i, status_col)
            c.fill = DONE if c.value == '제출' else TODO
    ws.freeze_panes = 'A4'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = '동인엔시스 | &A | 페이지 &P / 전체 &N'
    ws.oddFooter.center.size = 9


def write_checklist(project: Project) -> Path:
    data = project.data()
    biz = data.get('사업', {})
    wb = Workbook()
    ws = wb.active
    ws.title = '요청자료'
    tasks = {t.get('id'): t.get('이름', '') for t in data.get('과제목록') or []}
    rows = []
    for e in evidence_status(project):
        rows.append([e['id'], e.get('항목', ''), f"{e.get('관련과제', '')} {tasks.get(e.get('관련과제'), '')}".strip(),
                     e.get('종류', ''), '필수' if e.get('필수', True) else '선택', e.get('설명', ''),
                     e.get('담당', ''), e.get('마감', ''), '제출' if e['제출'] else '미제출',
                     '\n'.join(e['파일'][:6]) + (f"\n외 {len(e['파일']) - 6}개" if len(e['파일']) > 6 else '')])
    _sheet(ws, f"요청자료 체크리스트 — {biz.get('사업명', '')}",
           ['ID', '자료', '관련 과제', '종류', '구분', '요청 내용 (찍을 것·받을 것)', '담당', '마감', '상태', '받은 파일'],
           [7, 24, 22, 8, 7, 46, 10, 11, 9, 34], rows, status_col=9)
    ws.cell(ws.max_row + 2, 1, '파일 이름 규칙: 자료 ID로 시작 (예: E01_정면.jpg, E03_시험성적서.pdf) 또는 evidence/E01/ 폴더에 넣기').font = Font(name=F, size=9, color='5A6577')

    ws2 = wb.create_sheet('제출서류')
    rows2 = []
    ev_ok = {e['id']: e['제출'] for e in evidence_status(project)}
    for d in data.get('서류목록') or []:
        need = d.get('필요증빙') or []
        missing = [x for x in need if not ev_ok.get(x)]
        rows2.append([d.get('id', ''), d.get('서류명', ''), d.get('작성', ''), d.get('양식', '') or '-',
                      ', '.join(need) or '-', ', '.join(missing) or '없음', d.get('비고', '')])
    _sheet(ws2, '대금청구 제출서류', ['ID', '서류', '작성 방식', '양식 파일', '필요 자료', '빠진 자료', '비고'],
           [7, 26, 14, 28, 18, 18, 36], rows2)
    out = project.path('output', '요청자료_체크리스트.xlsx')
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
