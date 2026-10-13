"""기관 양식(DOCX·XLSX, HWP·HWPX 는 워드로 변환)의 빈칸 찾기(scan)와 채우기(fill). 결과물은 워드·엑셀(+PDF).

채우기 명세(fill/*.yaml):
  양식: 대금청구서.hwp                  # input/forms 기준 경로 (한글 양식은 work/converted/*.docx 로 바꿔 채움)
  출력: 01_대금청구서                   # output/ 아래 파일 이름(확장자 제외)
  채우기:                               # 라벨 기준: "라벨 > right" / "라벨 > down" / "라벨 > right > right"
    사업명 > right: "{{사업.사업명}}"
  셀:                                   # 주소 직접 지정 (표 번호·행·열은 scan 결과 기준, 0부터)
    - {표: 1, 행: 2, 열: 3, 값: "..."}
    - {시트: 청구서, 셀: D7, 값: "..."}  # XLSX
  바꾸기:                               # 문장 속 빈칸
    "년    월    일": "{{청구.청구일|날짜:%Y년 %m월 %d일}}"
  그림:
    - {파일: "evidence/E01_*", 라벨: "사진 > down", 폭mm: 70}   # 라벨 칸 안에 넣기
    - {파일: "evidence/E02_*", 위치: 끝, 폭mm: 120}             # 문서 끝에 붙이기
"""
from __future__ import annotations

import io
import json
import re
import warnings
from pathlib import Path

import yaml

from . import media
from .project import Project, render

warnings.filterwarnings('ignore', category=DeprecationWarning)

_DATE_BLANK = re.compile(r'(\d{4}|\s{2,})\s*년\s{2,}월\s{2,}일')


def norm(s: str) -> str:
    return re.sub(r'[\s:：*·ㆍ　]', '', s or '')


def _parse_path(path: str) -> tuple[str, list[str]]:
    parts = [p.strip() for p in path.split('>')]
    return parts[0], [p.lower() for p in parts[1:]] or ['right']


# ===================================================================== scan

def as_word(src: Path, work: Path) -> Path:
    """한글 양식이면 워드로 바꾼 파일 경로(work/converted/<이름>.docx), 아니면 그대로"""
    if src.suffix.lower() not in ('.hwp', '.hwpx'):
        return src
    from .hwp2docx import convert
    out = work / 'converted' / (src.stem + '.docx')
    if not out.exists() or out.stat().st_mtime < src.stat().st_mtime:
        out.parent.mkdir(parents=True, exist_ok=True)
        convert(src, out)
    return out


def scan(src: Path) -> dict:
    ext = src.suffix.lower()
    if ext == '.docx':
        tables, text = _scan_docx(src)
    elif ext in ('.xlsx', '.xlsm'):
        return _scan_xlsx(src)
    else:
        raise ValueError(f'양식 형식 미지원: {src.name} (DOCX·XLSX, 한글은 as_word 로 변환 후)')
    blanks = sorted({m.group(0) for m in _DATE_BLANK.finditer(text)})
    marks = sorted({m for m in re.findall(r'\((?:인|서명|직인)\)|\{\{[^}]+\}\}', text)})
    return {'tables': tables, 'date_blanks': blanks, 'marks': marks}


def _candidates(cells: dict[tuple[int, int], str]) -> list[str]:
    """라벨처럼 보이는 칸 → 오른쪽/아래가 빈칸이면 후보 경로"""
    out = []
    rows = sorted({r for r, _ in cells})
    for (r, c), t in sorted(cells.items()):
        if not t or len(t) > 30 or '\n' in t:
            continue
        right = [cc for (rr, cc) in cells if rr == r and cc > c]
        if right and not cells[(r, min(right))]:
            out.append(f'{t} > right')
            continue
        below = [rr for rr in rows if rr > r and (rr, c) in cells]
        if below and not cells[(min(below), c)]:
            out.append(f'{t} > down')
    return out


def _docx_grid(table) -> dict[tuple[int, int], object]:
    grid = {}
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            grid[(r, c)] = cell
    return grid


def _scan_docx(src: Path):
    import docx
    d = docx.Document(str(src))
    tables = []
    for i, t in enumerate(d.tables):
        grid = _docx_grid(t)
        cells, seen = {}, set()
        for (r, c), cell in grid.items():
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells[(r, c)] = cell.text.strip()
        tables.append({'표': i, '크기': f'{len(t.rows)}x{len(t.columns)}',
                       '칸': {f'{r},{c}': v for (r, c), v in sorted(cells.items()) if v},
                       '후보': _candidates(cells)})
    text = '\n'.join(p.text for p in d.paragraphs)
    return tables, text


def _scan_xlsx(src: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    sheets = []
    for ws in wb.worksheets:
        cells = {(c.row, c.column): str(c.value).strip() for row in ws.iter_rows() for c in row if c.value is not None}
        cands = []
        for (r, c), t in sorted(cells.items()):
            if len(t) <= 30 and (r, c + 1) not in cells:
                cands.append(f'{t} > right')
        sheets.append({'시트': ws.title, '칸': {ws.cell(r, c).coordinate: v for (r, c), v in sorted(cells.items())},
                       '후보': cands})
    return {'sheets': sheets, 'date_blanks': [], 'marks': []}


def write_scan(src: Path, result: dict, md_out: Path, yaml_out: Path, rel_form: str):
    lines = [f'# 양식 분석: {src.name}', '']
    paths = []
    for t in result.get('tables', []):
        lines.append(f"## 표 {t['표']} ({t['크기']}) {t.get('앞문단', '')}")
        for k, v in t['칸'].items():
            lines.append(f'- [{k}] {v}')
        if t['후보']:
            lines.append('- 빈칸 후보: ' + ', '.join(t['후보']))
            paths += t['후보']
        lines.append('')
    for s in result.get('sheets', []):
        lines.append(f"## 시트 {s['시트']}")
        lines += [f'- {k}: {v}' for k, v in s['칸'].items()]
        paths += s['후보']
    if result.get('date_blanks'):
        lines.append('## 날짜 빈칸: ' + ' / '.join(result['date_blanks']))
    if result.get('marks'):
        lines.append('## 표시: ' + ' / '.join(result['marks']))
    md_out.write_text('\n'.join(lines), encoding='utf-8')

    if not yaml_out.exists():                       # never overwrite a spec someone already filled in
        uniq = list(dict.fromkeys(paths))
        spec = {'양식': rel_form, '출력': src.stem,
                '채우기': {p: '' for p in uniq},
                '바꾸기': {b: '' for b in result.get('date_blanks', [])}}
        yaml_out.write_text('# scan 이 만든 초안. 값을 채우고 쓰지 않는 줄은 지운다.\n'
                            + yaml.safe_dump(spec, allow_unicode=True, sort_keys=False), encoding='utf-8')


# ===================================================================== fill

class FillLog:
    def __init__(self):
        self.applied, self.failed = [], []

    def ok(self, what, where=''):
        self.applied.append({'항목': what, '위치': str(where)})

    def fail(self, what, why):
        self.failed.append({'항목': what, '이유': why})


def fill(project: Project, spec_path: Path, pdf: bool = True) -> dict:
    spec = yaml.safe_load(spec_path.read_text(encoding='utf-8')) or {}
    data = project.data()
    src = project.path('input', 'forms', spec['양식']) if not Path(spec['양식']).parts[0] in ('input', 'docs') else project.path(spec['양식'])
    if not src.exists():
        raise SystemExit(f'양식 파일이 없습니다: {src}')
    out_base = project.path('output', spec.get('출력') or src.stem)
    out_base.parent.mkdir(parents=True, exist_ok=True)

    values = {k: render(v, data) for k, v in (spec.get('채우기') or {}).items() if v not in (None, '')}
    cells = [render(c, data) for c in (spec.get('셀') or [])]
    replaces = {k: render(v, data) for k, v in (spec.get('바꾸기') or {}).items() if v not in (None, '')}
    pics = spec.get('그림') or []
    log = FillLog()

    src = as_word(src, project.path('work'))
    ext = src.suffix.lower()
    if ext == '.docx':
        outputs = _fill_docx(src, out_base, values, cells, replaces, pics, project, log)
    elif ext in ('.xlsx', '.xlsm'):
        outputs = _fill_xlsx(src, out_base, values, cells, replaces, pics, project, log)
    else:
        raise SystemExit(f'양식 형식 미지원: {src.name}')

    if pdf:
        for o in list(outputs):
            if o.suffix in ('.docx', '.xlsx'):
                p = media.to_pdf(o, o.parent)
                if p:
                    outputs.append(p)

    leftovers = [str(v) for v in list(values.values()) + list(replaces.values()) + [c.get('값', '') for c in cells]
                 if '【확인필요' in str(v)]
    result = {'양식': str(src.name), '출력': [str(o.relative_to(project.dir)) for o in outputs],
              '적용': log.applied, '실패': log.failed, '확인필요값': leftovers}
    logs = project.path('work', 'logs')
    logs.mkdir(parents=True, exist_ok=True)
    (logs / f'{spec_path.stem}.fill.json').write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding='utf-8')
    return result


def _pictures(project: Project, pic: dict):
    files = project.resolve_files(pic['파일'])
    return [f for f in files if media.is_image(f)]


# ---- DOCX ------------------------------------------------------------------------------------

def _set_cell_text(cell, value: str):
    lines = str(value).split('\n')
    paras = cell.paragraphs
    first = paras[0]
    _set_para_text(first, lines[0])
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    for line in lines[1:]:
        np_ = cell.add_paragraph(line, style=first.style)
        np_.paragraph_format.alignment = first.paragraph_format.alignment


def _set_para_text(p, text: str):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def _docx_find(doc, label: str, dirs: list[str]):
    want = norm(label)
    for ti, t in enumerate(doc.tables):
        grid = _docx_grid(t)
        for (r, c), cell in grid.items():
            if norm(cell.text) != want:
                continue
            cur_r, cur_c, cur = r, c, cell
            ok = True
            for d in dirs:
                dr, dc = (0, 1) if d.startswith('r') else (1, 0) if d.startswith('d') else (0, -1) if d.startswith('l') else (-1, 0)
                nr, nc = cur_r + dr, cur_c + dc
                while (nr, nc) in grid and grid[(nr, nc)]._tc is cur._tc:   # step past merged span
                    nr, nc = nr + dr, nc + dc
                if (nr, nc) not in grid:
                    ok = False
                    break
                cur_r, cur_c, cur = nr, nc, grid[(nr, nc)]
            if ok:
                return ti, cur_r, cur_c, cur
    return None


def _docx_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs


def _fill_docx(src, out_base, values, cells, replaces, pics, project, log):
    import docx
    from docx.shared import Mm
    doc = docx.Document(str(src))
    for path, v in values.items():
        label, dirs = _parse_path(path)
        hit = _docx_find(doc, label, dirs)
        if not hit:
            log.fail(path, '라벨을 찾지 못함')
            continue
        ti, r, c, cell = hit
        _set_cell_text(cell, v)
        log.ok(path, f'표{ti}({r},{c})')
    for c in cells:
        try:
            _set_cell_text(doc.tables[int(c['표'])].cell(int(c['행']), int(c['열'])), c['값'])
            log.ok(f"표{c['표']}({c['행']},{c['열']})")
        except Exception as e:
            log.fail(f"표{c.get('표')}({c.get('행')},{c.get('열')})", str(e))
    for k, v in replaces.items():
        n = 0
        for p in _docx_paragraphs(doc):
            if k in p.text:
                _set_para_text(p, p.text.replace(k, str(v)))
                n += 1
        (log.ok if n else log.fail)(f'바꾸기 "{k}"', f'{n}곳' if n else '문장을 찾지 못함')
    seen_cells = set()
    for pic in pics:
        files = _pictures(project, pic)
        if not files:
            log.fail(f"그림 {pic['파일']}", '파일 없음')
            continue
        width = Mm(float(pic.get('폭mm', 70)))
        cell = None
        if pic.get('라벨'):
            hit = _docx_find(doc, *_parse_path(pic['라벨']))
            cell = hit[3] if hit else None
            if not hit:
                log.fail(f"그림 위치 {pic['라벨']}", '라벨을 찾지 못해 문서 끝에 붙임')
        for f in files:
            data, _ = media.prepare(f)
            if cell is not None:
                if id(cell._tc) not in seen_cells and not cell.text.strip():
                    para = cell.paragraphs[0]
                    seen_cells.add(id(cell._tc))
                else:
                    para = cell.add_paragraph()
                para.add_run().add_picture(io.BytesIO(data), width=width)
                log.ok(f'그림 {f.name}', pic.get('라벨'))
            else:
                doc.add_picture(io.BytesIO(data), width=width)
                log.ok(f'그림 {f.name}', '문서 끝')
    out = Path(str(out_base) + '.docx')
    doc.save(str(out))
    return [out]


# ---- XLSX ------------------------------------------------------------------------------------

def _xl_set(cell, v):
    cell.value = v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and cell.number_format == 'General':
        cell.number_format = '#,##0'


def _fill_xlsx(src, out_base, values, cells, replaces, pics, project, log):
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.load_workbook(str(src))

    def merged_range(ws, cell):
        for rng in ws.merged_cells.ranges:
            if cell.coordinate in rng:
                return rng
        return None

    def step(ws, cell, d):
        rng = merged_range(ws, cell)
        r, c = cell.row, cell.column
        if d.startswith('r'):
            c = (rng.max_col if rng else c) + 1
        elif d.startswith('d'):
            r = (rng.max_row if rng else r) + 1
        nxt = ws.cell(r, c)
        rng2 = merged_range(ws, nxt)
        return ws.cell(rng2.min_row, rng2.min_col) if rng2 else nxt

    def find(label, dirs):
        want = norm(label)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and norm(str(cell.value)) == want:
                        cur = cell
                        for d in dirs:
                            cur = step(ws, cur, d)
                        return ws, cur
        return None

    for path, v in values.items():
        hit = find(*_parse_path(path))
        if not hit:
            log.fail(path, '라벨을 찾지 못함')
            continue
        ws, cell = hit
        _xl_set(cell, v)
        log.ok(path, f'{ws.title}!{cell.coordinate}')
    for c in cells:
        try:
            ws = wb[c['시트']] if c.get('시트') else wb.active
            _xl_set(ws[c['셀']], c['값'])
            log.ok(f"{ws.title}!{c['셀']}")
        except Exception as e:
            log.fail(str(c), str(e))
    for k, v in replaces.items():
        n = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and k in cell.value:
                        cell.value = cell.value.replace(k, str(v))
                        n += 1
        (log.ok if n else log.fail)(f'바꾸기 "{k}"', f'{n}곳' if n else '문장을 찾지 못함')
    tmp = []
    for pic in pics:
        hit = find(*_parse_path(pic['라벨'])) if pic.get('라벨') else None
        ws, cell = hit if hit else (wb.active, None)
        for f in _pictures(project, pic):
            data, (w, h) = media.prepare(f, 1200)
            img = XLImage(io.BytesIO(data))
            px = float(pic.get('폭mm', 70)) / 25.4 * 96
            img.width, img.height = px, px * h / w
            ws.add_image(img, cell.coordinate if cell else 'A1')
            log.ok(f'그림 {f.name}', f'{ws.title}!{cell.coordinate if cell else "A1"}')
    out = Path(str(out_base) + '.xlsx')
    wb.save(str(out))
    return [out]
