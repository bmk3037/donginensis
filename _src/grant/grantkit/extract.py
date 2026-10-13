"""PDF·DOCX·HWP·HWPX·XLSX → 마크다운 텍스트 (Claude 가 읽고 project.yaml 을 채우는 재료)."""
from __future__ import annotations

import subprocess
import warnings
from pathlib import Path

DOC_EXT = {'.pdf', '.docx', '.hwp', '.hwpx', '.xlsx', '.xlsm', '.txt', '.md', '.csv'}


def to_markdown(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == '.pdf':
        return _pdf(path)
    if ext == '.docx':
        return _docx(path)
    if ext in ('.hwp', '.hwpx'):
        return _hwp(path)
    if ext in ('.xlsx', '.xlsm'):
        return _xlsx(path)
    if ext in ('.txt', '.md', '.csv'):
        return path.read_text(encoding='utf-8', errors='replace')
    raise ValueError(f'지원하지 않는 형식: {path.name}')


def _pdf(path: Path) -> str:
    import fitz
    out = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, 1):
            out.append(f'\n\n<!-- page {i} -->\n')
            out.append(page.get_text('text'))
            try:
                for t in page.find_tables().tables:
                    md = t.to_markdown()
                    if md.strip():
                        out.append('\n[표]\n' + md)
            except Exception:  # table detection is best-effort
                pass
    text = ''.join(out)
    if len(text.strip()) < 50 * max(1, text.count('<!-- page')):
        text += '\n\n<!-- 텍스트가 거의 없음: 스캔 PDF 일 수 있음. 원본 파일(HWP·DOCX)이나 OCR 이 필요 -->'
    return text


def _cell_text(cell) -> str:
    return ' '.join(p.text.strip() for p in cell.paragraphs if p.text.strip()).replace('|', '/')


def _docx_table(table) -> str:
    rows = []
    for r in table.rows:
        seen, cells = set(), []
        for c in r.cells:
            if id(c._tc) in seen:            # merged cell repeats
                continue
            seen.add(id(c._tc))
            cells.append(_cell_text(c))
        rows.append(cells)
    if not rows:
        return ''
    lines = ['| ' + ' | '.join(r) + ' |' for r in rows]
    lines.insert(1, '|' + ' --- |' * len(rows[0]))
    return '\n'.join(lines)


def _docx(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split('}')[-1]
        if tag == 'p':
            p = Paragraph(el, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or '').lower() if p.style is not None else ''
            if style.startswith('heading') or style.startswith('제목'):
                level = ''.join(ch for ch in style if ch.isdigit()) or '1'
                out.append('#' * min(int(level) + 1, 6) + ' ' + t)
            else:
                out.append(t)
        elif tag == 'tbl':
            out.append(_docx_table(Table(el, d)))
    return '\n\n'.join(out)


def _hwp(path: Path) -> str:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            from hwpx import HwpxDocument
            return HwpxDocument.open(str(path)).export_markdown()
    except Exception as e:
        if path.suffix.lower() == '.hwp':     # independent HWP5 reader as a fallback
            r = subprocess.run(['hwp5txt', str(path)], capture_output=True, text=True)
            if r.returncode == 0 and r.stdout.strip():
                return r.stdout + '\n\n<!-- hwp5txt 로 추출: 표 내용은 빠질 수 있음 -->'
        raise RuntimeError(f'{path.name} 을 읽지 못했습니다: {e}')


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f'## 시트: {ws.title}')
        for row in ws.iter_rows():
            vals = [('' if c.value is None else str(c.value).replace('\n', ' ')) for c in row]
            if any(vals):
                coord = row[0].coordinate
                out.append(f'{coord}: ' + ' | '.join(vals).rstrip(' |'))
    return '\n'.join(out)


def extract_all(src_dirs: list[Path], out_dir: Path) -> list[tuple[Path, Path | None, str]]:
    """returns [(source, written_md or None, message)]"""
    out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for d in src_dirs:
        if not d.exists():
            continue
        for f in sorted(d.rglob('*')):
            if not f.is_file() or f.name.startswith(('.', '~$')):
                continue
            if f.suffix.lower() not in DOC_EXT:
                results.append((f, None, '건너뜀(문서 아님)'))
                continue
            rel = f.relative_to(d.parent)
            target = out_dir / (str(rel).replace('/', '__') + '.md')
            try:
                md = to_markdown(f)
                target.write_text(f'<!-- source: {rel} -->\n' + md, encoding='utf-8')
                results.append((f, target, f'{len(md):,}자'))
            except Exception as e:
                results.append((f, None, f'실패: {e}'))
    return results
