"""통합 기록양식집(DI-IMS-F-01) 엑셀 생성기.

양식 내용은 forms_data.py에 있고, 이 파일은 공통 배치(머리 입력란, 표, 결재란, 인쇄 설정)를 만든다.

    python3 _src/iso/build_forms.py            # files/iso/DI-IMS-F-01_Rev00.xlsx 생성
    python3 _src/iso/build_forms.py out.xlsx   # 다른 경로로 생성
"""
import itertools
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

import forms_data as D

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, 'files', 'iso', 'DI-IMS-F-01_Rev00.xlsx')

FONT = 'Malgun Gothic'
BLUE = '1857A5'
INK = '263445'
LINE = '8A9BB0'        # 흑백 인쇄에서도 보이는 테두리
LABEL_FILL = 'E8EEF6'  # 머리 입력란 항목명
PRE_FILL = 'F4F6F9'    # 미리 적힌 기준(입력 칸 아님)
WHITE = 'FFFFFF'

SIDE = Side(style='thin', color=LINE)
BOX = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
NEG_FILL = PatternFill('solid', fgColor='F8D7DA')
WARN_FILL = PatternFill('solid', fgColor='FFF3CD')

NEGATIVE = ['불합격', '보류', '부적합', '미흡', '미준수', '미확인', '실패', '사용금지', '중지·개선',
            '재교육·동행', 'C 발주중지', '중요']

# 표 전체 너비(엑셀 열 너비 단위): 열 너비 비율은 유지하고 용지 폭에 맞게 늘린다
TARGET_W = {'landscape': 140, 'portrait': 95}
# 인쇄 영역 높이(pt): A4에서 위·아래 여백을 뺀 값
PRINT_H = {'landscape': 505, 'portrait': 750}


def fill(color):
    return PatternFill('solid', fgColor=color)


def units(text, size=10):
    """셀 너비 단위로 본 글자 폭 추정 (엑셀 맑은 고딕 10pt 기준: 한글 1.7, 영숫자 1.0)."""
    total = 0.0
    for ch in str(text):
        if ord(ch) >= 0x1100:
            total += 1.7
        elif ch == ' ':
            total += 0.6
        else:
            total += 1.0
    return total * size / 10


def n_lines(text, width, size=10):
    if not text:
        return 1
    usable = max(width - 1.5, 3)
    return sum(max(1, math.ceil(units(part, size) / usable)) for part in str(text).split('\n'))


def text_height(text, width, size=10, minimum=18, pad=6):
    return max(minimum, n_lines(text, width, size) * size * 1.35 + pad)


class Sheet:
    def __init__(self, wb, spec, logo):
        self.spec = spec
        self.ws = wb.create_sheet(spec['sheet'])
        raw = [w for _, w in spec['cols']]
        k = TARGET_W[spec.get('orient', 'landscape')] / float(sum(raw))
        self.widths = [round(w * k, 1) for w in raw]
        self.n = len(self.widths)
        self.row = 1
        self.heights = {}
        self.logo = logo
        ws = self.ws
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(self.widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---------- 기본 도구 ----------
    def width(self, c1, c2):
        return sum(self.widths[c1 - 1:c2])

    def set_h(self, r, h):
        self.heights[r] = h
        self.ws.row_dimensions[r].height = h

    def cell(self, r, c1, c2=None, value=None, size=10, bold=False, color=INK, bg=None,
             border=True, h='left', v='center', wrap=True, fmt=None):
        c2 = c2 or c1
        ws = self.ws
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        for c in range(c1, c2 + 1):
            x = ws.cell(r, c)
            x.font = Font(name=FONT, size=size, bold=bold, color=color)
            x.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
            if bg:
                x.fill = fill(bg)
            if border:
                x.border = BOX
            if fmt:
                x.number_format = fmt
        if value is not None:
            ws.cell(r, c1).value = value
        return ws.cell(r, c1)

    def dv_list(self, ref, options, prompt=None):
        dv = DataValidation(type='list', formula1='"%s"' % ','.join(options), allow_blank=True)
        dv.error = '목록에서 고르십시오: ' + ', '.join(options)
        dv.errorStyle = 'warning'
        if prompt:
            dv.prompt = prompt
        self.ws.add_data_validation(dv)
        dv.add(ref)

    def dv_whole(self, ref, lo, hi):
        dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi),
                            allow_blank=True)
        dv.error = '%d~%d 사이 정수를 입력하십시오.' % (lo, hi)
        self.ws.add_data_validation(dv)
        dv.add(ref)

    # ---------- 머리 ----------
    def header(self):
        s = self.spec
        ws = self.ws
        self.set_h(1, 22)
        if self.logo:
            img = XLImage(self.logo)
            img.width, img.height = 95, 18
            ws.add_image(img, 'A1')
        self.set_h(2, 28)
        self.cell(2, 1, self.n, s['title'], size=16, bold=True, color=BLUE, border=False)
        line = D.doc_line(s['code'])
        if s.get('owner'):
            line += '   ·   담당: %s | 작성: %s' % (s['owner'], s['when'])
        self.set_h(3, text_height(line, sum(self.widths), 9, 16, 2))
        self.cell(3, 1, self.n, line, size=9, border=False)
        self.row = 4

    def _span_for(self, c1, c2, label):
        """묶음 [c1, c2]에서 항목명 칸 끝 열을 정한다(나머지는 입력 칸)."""
        need = min(units(label, 9) + 1.5, 16)
        end = c1
        while self.width(c1, end) < need and end < c2 - 1:
            end += 1
        return end

    def _score(self, groups, labels):
        total = 0.0
        widths = []
        for (c1, c2), label in zip(groups, labels):
            le = self._span_for(c1, c2, label)
            lw, vw = self.width(c1, le), self.width(le + 1, c2)
            widths.append(lw + vw)
            total += max(0, 18 - vw) * 3                       # 입력 칸이 좁으면 감점
            total += max(0, lw - (units(label, 9) + 8)) * 1     # 항목명 칸이 지나치게 넓으면 감점
            total += (n_lines(label, lw, 9) - 1) * 4            # 항목명 줄바꿈 감점
        avg = sum(widths) / len(widths)
        total += sum(abs(w - avg) for w in widths) * 0.3
        return total

    def _groups(self, labels):
        """항목 수만큼 열을 묶는다(각 묶음 = 항목명 칸 + 입력 칸)."""
        k = len(labels)
        if k == 1:
            return [(1, self.n)]
        best, best_g = None, None
        n = self.n
        for cuts in itertools.combinations(range(2, n), k - 1):
            bounds = [1] + list(cuts) + [n + 1]
            groups = [(bounds[i], bounds[i + 1] - 1) for i in range(k)]
            if any(b - a < 1 for a, b in groups):
                continue
            sc = self._score(groups, labels)
            if best is None or sc < best:
                best, best_g = sc, groups
        return best_g

    def _packed_fields(self):
        """두 칸짜리 줄은 모아서 가로 A4는 한 줄에 3개, 세로는 2개씩 다시 배치한다."""
        per = 3 if self.spec.get('orient', 'landscape') == 'landscape' and self.n >= 6 else 2
        if per == 3:
            # 3개씩 놓으면 입력 칸이 너무 좁아지는 열 구성이면 2개씩 놓는다
            sample = ['프로젝트 번호'] * 3
            g = self._groups(sample)
            narrow = min(self.width(self._span_for(c1, c2, sample[0]) + 1, c2) for c1, c2 in g)
            if narrow < 17:
                per = 2
        out, buf = [], []

        def flush():
            while buf:
                out.append(buf[:per])
                del buf[:per]

        for frow in self.spec.get('fields', []):
            if len(frow) == 1:
                flush()
                out.append(frow)
            else:
                buf.extend(frow)
        flush()
        return out

    def fields(self):
        rows = self._packed_fields()
        if not rows:
            return
        self.set_h(self.row, 4)
        self.row += 1
        for frow in rows:
            groups = self._groups([f[0] for f in frow])
            r = self.row
            h = 20
            for (c1, c2), f in zip(groups, frow):
                label = f[0]
                le = self._span_for(c1, c2, label)
                self.cell(r, c1, le, label, size=9, bg=LABEL_FILL)
                val = f[2] if len(f) > 2 else None
                self.cell(r, le + 1, c2, val, size=10, bg=WHITE)
                if len(f) > 1 and f[1]:
                    self.dv_list('%s%d' % (get_column_letter(le + 1), r), f[1])
                h = max(h, text_height(label, self.width(c1, le), 9, 20, 4))
            self.set_h(r, h)
            self.row += 1

    def notes(self):
        ns = self.spec.get('notes', [])
        if not ns:
            return
        self.set_h(self.row, 5)
        self.row += 1
        for t in ns:
            self.cell(self.row, 1, self.n, t, size=9, border=False)
            self.set_h(self.row, text_height(t, sum(self.widths), 9, 14, 2))
            self.row += 1

    # ---------- 표 ----------
    def table(self):
        s = self.spec
        ws = self.ws
        self.set_h(self.row, 5)
        self.row += 1
        hr = self.row
        self.header_row = hr
        h = 26
        for i, (name, w) in enumerate(s['cols'], 1):
            self.cell(hr, i, None, name, size=9.5, bold=True, color=WHITE, bg=BLUE, h='center')
            h = max(h, text_height(name, w, 9.5, 26, 6))
        self.set_h(hr, h)
        first = hr + 1
        base_h = s.get('row_h', 24)
        rows = list(s.get('rows', [])) + [[]] * self._blank_rows(h, base_h)
        for k, data in enumerate(rows):
            r = first + k
            rh = base_h
            for i in range(1, self.n + 1):
                v = data[i - 1] if i - 1 < len(data) else None
                if v == '':
                    v = None
                pre = v is not None and not str(v).startswith('=')
                self.cell(r, i, None, v, size=9.5, bg=PRE_FILL if pre else WHITE,
                          h='center' if self.widths[i - 1] <= 7 else 'left')
                if pre:
                    rh = max(rh, text_height(v, self.widths[i - 1], 9.5, base_h, 6))
            self.set_h(r, rh)
        last = first + len(rows) - 1
        self.first, self.last = first, last
        # 수식·검증
        for col, formula in s.get('formulas', {}).items():
            for r in range(first, last + 1):
                c = ws.cell(r, col)
                c.value = formula.format(r=r)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                c.fill = fill(PRE_FILL)
        for col, fmt in s.get('formats', {}).items():
            for r in range(first, last + 1):
                ws.cell(r, col).number_format = fmt
        for col, opts in s.get('dv', {}).items():
            L = get_column_letter(col)
            ref = '%s%d:%s%d' % (L, first, L, last)
            if isinstance(opts, tuple):
                self.dv_whole(ref, *opts)
            else:
                self.dv_list(ref, opts)
            for r in range(first, last + 1):
                ws.cell(r, col).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        # 부정 판정은 붉게
        for col in list(s.get('dv', {}).keys()) + list(s.get('formulas', {}).keys()):
            L = get_column_letter(col)
            ref = '%s%d:%s%d' % (L, first, L, last)
            cond = 'OR(%s)' % ','.join('%s%d="%s"' % (L, first, w) for w in NEGATIVE)
            ws.conditional_formatting.add(ref, FormulaRule(formula=[cond], fill=NEG_FILL))
        for col, (red, yellow) in s.get('risk_cf', {}).items():
            L = get_column_letter(col)
            ref = '%s%d:%s%d' % (L, first, L, last)
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=['AND(ISNUMBER(%s%d),%s%d>=%s)' % (L, first, L, first, red)], fill=NEG_FILL))
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=['AND(ISNUMBER(%s%d),%s%d>=%s)' % (L, first, L, first, yellow)], fill=WARN_FILL))
        for col in s.get('due_cf', []):
            L = get_column_letter(col)
            ref = '%s%d:%s%d' % (L, first, L, last)
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=['AND(ISNUMBER(%s%d),%s%d<TODAY())' % (L, first, L, first)], fill=NEG_FILL))
            ws.conditional_formatting.add(ref, FormulaRule(
                formula=['AND(ISNUMBER(%s%d),%s%d<TODAY()+30)' % (L, first, L, first)], fill=WARN_FILL))
        self.row = last + 1
        if s.get('after_table'):
            s['after_table'](self)

    def _row_height(self, data, base_h):
        rh = base_h
        for i, v in enumerate(data):
            if v not in (None, '') and not str(v).startswith('='):
                rh = max(rh, text_height(v, self.widths[i], 9.5, base_h, 6))
        return rh

    def _blank_rows(self, header_h, base_h):
        """대장 양식은 빈 행을 마지막 인쇄 쪽 끝까지 채운다(최소 spec['blank']행)."""
        s = self.spec
        minimum = s.get('blank', 0)
        if s.get('mode', 'single') != 'register':
            return minimum
        cap = PRINT_H[s.get('orient', 'landscape')]
        top = sum(self.heights.values())          # 머리·입력란·안내·표 머리글까지
        pre = sum(self._row_height(d, base_h) for d in s.get('rows', []))
        tail = 0 if s.get('sign', ('',)) is None else 6 + 16 + 28 + 16
        k = 1
        while True:
            room = k * cap - (k - 1) * header_h - top - pre - tail
            n = int(room // base_h)
            if n >= minimum:
                return max(n - 1, minimum)
            k += 1

    # ---------- 결재란 ----------
    def sign(self):
        labels = self.spec.get('sign', ('작성', '검토', '승인'))
        if labels is None:
            return
        self.set_h(self.row, 6)
        self.row += 1
        r = self.row
        # 오른쪽부터 열을 모아 4칸(구분 + 결재 3칸)을 만든다
        groups, c, need = [], self.n, 4
        while len(groups) < need and c >= 1:
            end = c
            while self.width(c, end) < 10 and c > 1:
                c -= 1
            groups.append((c, end))
            c -= 1
        groups.reverse()
        left_end = groups[0][0] - 1
        attach = '첨부·대체 기록 번호와 보관 경로 / 적용 제외(해당없음) 사유 / 미결 확인 :'
        if left_end >= 1 and self.width(1, left_end) < 24:
            # 왼쪽 칸이 좁으면 첨부란을 결재란 위 한 줄로 둔다
            self.cell(r, 1, self.n, attach, size=9, bg=WHITE, v='top')
            self.set_h(r, 30)
            r += 1
            groups[0] = (1, groups[0][1])
            left_end = 0
        heads = ['구분'] + list(labels)
        rows = [('', heads, 16), ('성명·서명', None, 28), ('일자', None, 16)]
        for k, (title, values, h) in enumerate(rows):
            rr = r + k
            self.set_h(rr, h)
            for j, (c1, c2) in enumerate(groups):
                if k == 0:
                    self.cell(rr, c1, c2, heads[j], size=9, bold=True, bg=LABEL_FILL, h='center')
                elif j == 0:
                    self.cell(rr, c1, c2, title, size=9, bg=LABEL_FILL, h='center')
                else:
                    self.cell(rr, c1, c2, None, size=10, bg=WHITE)
        if left_end >= 1:
            ws = self.ws
            ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=left_end)
            for rr in range(r, r + 3):
                for cc in range(1, left_end + 1):
                    x = ws.cell(rr, cc)
                    x.border = BOX
                    x.fill = fill(WHITE)
            x = ws.cell(r, 1)
            x.value = attach
            x.font = Font(name=FONT, size=9, color=INK)
            x.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        self.row = r + 3

    # ---------- 인쇄 ----------
    def page(self):
        s = self.spec
        ws = self.ws
        orient = s.get('orient', 'landscape')
        ws.page_setup.orientation = orient
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        mode = s.get('mode', 'single')
        if mode == 'single':
            ws.page_setup.fitToHeight = 1
        else:
            ws.page_setup.fitToHeight = 0
            if getattr(self, 'header_row', None):
                ws.print_title_rows = '%d:%d' % (self.header_row, self.header_row)
                ws.freeze_panes = 'A%d' % (self.header_row + 1)
        ws.page_margins.left = ws.page_margins.right = 0.45
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.55
        ws.page_margins.header = 0.25
        ws.page_margins.footer = 0.25
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = '동인엔시스 | %s Rev.00 | 페이지 &P / 전체 &N' % D.doc_no(s['code'])
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.font = FONT
        last_col = get_column_letter(self.n)
        ws.print_area = 'A1:%s%d' % (last_col, self.row - 1)

    def build(self):
        self.header()
        self.fields()
        self.notes()
        if self.spec.get('cols_table', True) and self.spec.get('rows') is not None:
            self.table()
        if self.spec.get('custom'):
            self.spec['custom'](self)
        self.sign()
        self.page()
        return self

    def total_height(self):
        return sum(self.heights.values())


def build(out=OUT):
    wb = Workbook()
    wb.remove(wb.active)
    logo = os.path.join(HERE, 'logo.png')
    report = []
    for spec in D.sheets():
        sh = Sheet(wb, spec, logo).build()
        mode = spec.get('mode', 'single')
        if mode == 'single':
            scale = PRINT_H[spec.get('orient', 'landscape')] / sh.total_height()
            report.append('%-9s single  height %4dpt  scale~%3d%%' % (spec['sheet'], sh.total_height(), min(100, scale * 100)))
        else:
            report.append('%-9s register height %4dpt' % (spec['sheet'], sh.total_height()))
    wb.properties.creator = '동인엔시스'
    wb.properties.title = '통합 기록양식집 DI-IMS-F-01 Rev.00'
    wb.save(out)
    print('\n'.join(report))
    print('saved', out)


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else OUT)
